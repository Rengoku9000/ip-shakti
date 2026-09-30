from pptx import Presentation

prs = Presentation(r'c:\Projects\drugvista\IP-SAKTI_Sahayak_SIH2026.pptx')
for i, slide in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} ===')
    for s in slide.shapes:
        if s.top < 1200000 or s.top > 6200000:
            text = s.text_frame.text[:40].strip() if s.has_text_frame else ''
            print(f'  {s.name} (type {s.shape_type}): top={s.top}, left={s.left}, text="{text}"')
