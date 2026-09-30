from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

pptx_path = r"c:\Projects\drugvista\IP-SAKTI_Sahayak_SIH2026.pptx"
prs = Presentation(pptx_path)

titles = [
    (1, "PROPOSED SOLUTION"),
    (2, "TECHNICAL APPROACH & ARCHITECTURE"),
    (3, "FEASIBILITY, VIABILITY & ROADMAP"),
    (4, "IMPACT, BENEFITS & STAKEHOLDERS"),
    (5, "RESEARCH, REFERENCES & PROTOTYPE"),
]

for idx, title_text in titles:
    slide = prs.slides[idx]
    for s in slide.shapes:
        if s.name == 'Title 1':
            s.text_frame.text = title_text
            p = s.text_frame.paragraphs[0]
            p.font.name = "Segoe UI"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(11, 37, 69)
            print(f"Slide {idx+1} Title set to: {title_text}")
        elif 'oval' in s.name.lower() or (s.top < 500000 and s.left < 1000000 and s.width < 2500000):
            tf = s.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.text = "OUTLAWS"
            p.font.name = "Segoe UI"
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = RGBColor(11, 37, 69)
            p.alignment = PP_ALIGN.CENTER
            print(f"Slide {idx+1} Oval badge set to: OUTLAWS")

prs.save(pptx_path)
print("Saved cleanly!")
