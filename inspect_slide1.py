from pptx import Presentation

p = r"C:\Projects\Secure mesh\docs\PPT\SecureMesh.pptx"
prs = Presentation(p)

def inspect_slide(s_idx):
    slide = prs.slides[s_idx]
    print(f"=== SLIDE {s_idx + 1} ===")
    for sh in slide.shapes:
        if sh.has_text_frame:
            full = " ".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())
            if full:
                print(f"[{sh.name}]: {full[:80]}")
                for p_idx, p in enumerate(sh.text_frame.paragraphs):
                    for r_idx, r in enumerate(p.runs):
                        print(f"   P{p_idx} R{r_idx}: '{r.text}' (bold={r.font.bold}, size={r.font.size})")

inspect_slide(0)
