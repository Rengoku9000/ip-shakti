from pptx import Presentation

p = r"C:\Projects\Secure mesh\docs\PPT\SecureMesh.pptx"
prs = Presentation(p)

with open("slide_shapes_dump.txt", "w", encoding="utf-8") as out:
    for s_idx, slide in enumerate(prs.slides):
        out.write(f"\n============================== SLIDE {s_idx + 1} ==============================\n")
        for sh in slide.shapes:
            txt = ""
            if sh.has_text_frame:
                txt = " | ".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())
            out.write(f"{sh.name} [{sh.shape_type}]: \"{txt}\"\n")

print("Dumped successfully with UTF-8")
