from pptx import Presentation

p = r"C:\Projects\Secure mesh\docs\PPT\SecureMesh.pptx"
prs = Presentation(p)

with open("all_runs_dump.txt", "w", encoding="utf-8") as out:
    for s_idx, slide in enumerate(prs.slides):
        out.write(f"\n============================== SLIDE {s_idx + 1} ==============================\n")
        for sh in slide.shapes:
            if sh.has_text_frame:
                full = " ".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())
                if full:
                    out.write(f"[{sh.name}]: {full}\n")
                    for p_idx, p in enumerate(sh.text_frame.paragraphs):
                        for r_idx, r in enumerate(p.runs):
                            out.write(f"   P{p_idx} R{r_idx}: '{r.text}'\n")

print("Wrote all_runs_dump.txt")
