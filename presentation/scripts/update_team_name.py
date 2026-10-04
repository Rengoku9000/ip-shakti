from pptx import Presentation

pptx_path = r"c:\Projects\drugvista\presentation\source\IP-SAKTI_Sahayak_SIH2026.pptx"
prs = Presentation(pptx_path)

# 1. Slide 1 update
slide1 = prs.slides[0]
for shape in slide1.shapes:
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            if "Team Name" in p.text:
                for run in p.runs:
                    if "IP-SAKTI" in run.text:
                        run.text = run.text.replace("IP-SAKTI Team", "OUTLAWS").replace("IP-SAKTI", "OUTLAWS")
                if "IP-SAKTI" in p.text:
                    p.text = p.text.replace("IP-SAKTI Team", "OUTLAWS").replace("IP-SAKTI", "OUTLAWS")
                print(f"Updated Slide 1: {p.text}")

# 2. Slides 2 to 6 top-left oval badge update
for idx in range(1, len(prs.slides)):
    slide = prs.slides[idx]
    for s in slide.shapes:
        if s.top < 600000 and s.left < 1500000:
            if s.has_text_frame and "IP-SAKTI" in s.text_frame.text:
                s.text_frame.text = "OUTLAWS"
                p = s.text_frame.paragraphs[0]
                p.font.name = "Segoe UI"
                p.font.bold = True
                print(f"Updated Slide {idx+1} oval badge to OUTLAWS")

prs.save(pptx_path)
print("Saved presentation with OUTLAWS team name successfully.")
